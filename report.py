"""Word reports from persisted scan evidence; never implies manual validation."""
import re
from io import BytesIO
from report_data import summarize,chart_png,public_evidence,redact,test_inventory
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

def safe(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]', '', str(value))

def export(job, path):
    d = Document()
    for border in list(d.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    sec = d.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.75)
    d.styles['Normal'].font.name = 'Calibri'
    d.styles['Normal'].font.size = Pt(10)
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].font.color.rgb = RGBColor(0,0,0)
    d.add_heading('Pentest Lab Scan Report',0)
    d.add_paragraph('Automated assessment of the configured target and exact host scope. Tool findings require manual validation before they are treated as confirmed vulnerabilities.')
    table=d.add_table(rows=0, cols=2); table.style='Light Shading Accent 1'
    for key,value in [('Version',job.get('version','Earlier version')),('Target',job.get('target','')),('Job',job['id']),('Status',job['status']),('Started',job['created']),('Finished',job.get('finished','In progress')),('Scope',', '.join(job.get('scope',[])))]:
        cells=table.add_row().cells; cells[0].text=key; cells[1].text=safe(redact(value))
    summary=summarize(job)
    d.add_heading('Executive summary',1)
    d.add_paragraph(f"{summary['total']} unique automated vulnerability candidates and {len(summary['observations'])} recon observations. {summary['raw_candidates']-summary['total']} duplicate candidate entries consolidated. Confidence is Unverified.")
    d.add_paragraph(summary['notice'])
    d.add_picture(BytesIO(chart_png(summary)),width=Inches(6))
    for row in summary['distribution']:d.add_paragraph(f"{row['category']}: {row['count']} ({row['percent']:.1f}%)")
    ledger=test_inventory(job)
    d.add_page_break()
    d.add_heading('Tests performed',1)
    d.add_paragraph(f"{ledger['attempts']} recorded command attempts; {ledger['completed']} completed; {ledger['failed']} failed or interrupted; {ledger['running']} running.")
    d.add_paragraph(ledger['notice'])
    checks=d.add_table(rows=1,cols=5);checks.style='Light Shading Accent 1'
    for cell,text in zip(checks.rows[0].cells,['Tool and test','Attempts','OK','Failed','Candidates']):cell.text=text
    for row in ledger['rows']:
        cells=checks.add_row().cells
        for cell,text in zip(cells,[row['tool']+' — '+row['test']+'\nStatus: '+row['state'],row['attempts'],row['completed'],row['failed'],row['candidates']]):cell.text=safe(text)
    checks.autofit=False
    widths=[3.2,.7,.55,.65,.9]
    repeat=OxmlElement('w:tblHeader');checks.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in checks.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(9)
    for col,width in zip(checks.columns,widths):col.width=Inches(width)
    d.add_heading('Finding register',1)
    if not summary['findings']:d.add_paragraph('No vulnerability candidates recorded. Review incomplete coverage below.')
    for i,f in enumerate(summary['findings'],1):d.add_paragraph(safe(f"F-{i:03d} | {f['category']} | {f['title']} | Unverified | {redact(f.get('url',''))}"))
    d.add_heading('Coverage and limitations',1)
    d.add_paragraph('A skipped, failed, timed out or cancelled stage has not completed its checks. No findings does not establish that the target is secure. Archived URLs may be stale. Discovery and parameter results are observations, not vulnerability confirmations. Raw evidence is retained in the job folder.')
    d.add_paragraph('Configured limits: '+', '.join(f'{key} {value}' for key,value in job.get('limits',{}).items()))
    if job.get('discovery_counts'): d.add_paragraph('Discovery coverage: '+safe(job['discovery_counts']))
    if job.get('routing_limits'): d.add_paragraph('Routing coverage: '+safe(job['routing_limits']))
    if job.get('tools'):d.add_paragraph('Selected tools: '+', '.join(job['tools']))
    if job.get('disabled_tools'):d.add_paragraph('Tools not selected: '+', '.join(job['disabled_tools']))
    if job.get('passive_note'):d.add_paragraph(safe(job['passive_note']))
    for stage in job.get('stages',[]):
        d.add_paragraph(safe(stage['name']+' — '+stage['status']+': '+redact(stage.get('detail',''))))
    if job.get('coverage'):
        d.add_heading('Tool coverage',1)
        for name,item in job['coverage'].items():d.add_paragraph(safe(name+': '+item['state']+' '+item.get('detail','')))
    if job.get('summary'):d.add_paragraph('Result counts: '+safe(job['summary']))
    findings=summary['findings']+summary['observations']
    d.add_heading('Findings and observations',1)
    d.add_paragraph(f'{len(findings)} findings or observations were recorded.')
    for i,f in enumerate(findings,1):
        d.add_heading(safe(f"{'F' if f.get('kind')!='observation' else 'O'}-{i:03d} {f['title']}"),2)
        d.add_paragraph('Confidence: '+f.get('confidence','Observation')).paragraph_format.keep_with_next=True
        d.add_paragraph('Sources: '+', '.join(f.get('sources',[f.get('tool','')])) ).paragraph_format.keep_with_next=True
        for key in ['kind','severity','url','method','parameter','observed_at','recommendation']:
            if f.get(key):d.add_paragraph(safe(key.capitalize()+': '+redact(f[key]))).paragraph_format.keep_with_next=True
        d.add_paragraph(safe('Evidence: '+public_evidence(f)))
    footer=sec.footer.paragraphs[0]
    footer.alignment=1
    run=footer.add_run('Developed by Nowan');run.font.size=Pt(9);run.font.color.rgb=RGBColor(90,100,115)
    d.core_properties.author='Nowan'
    d.save(path)
